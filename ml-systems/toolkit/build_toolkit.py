"""Builds ml_capacity_toolkit.xlsx from prices.json. Run: python build_toolkit.py"""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

P = json.load(open("prices.json"))
wb = Workbook()
F = lambda **k: Font(name="Arial", size=10, **k)
H = Font(name="Arial", size=11, bold=True, color="002D5C")
T = Font(name="Arial", size=14, bold=True, color="002D5C")
INP = Font(name="Arial", size=10, color="0000FF")
YEL = PatternFill("solid", fgColor="FFF2CC")
GREY = PatternFill("solid", fgColor="F2F2F2")

def title(ws, t, sub):
    ws["A1"] = t; ws["A1"].font = T
    ws["A2"] = sub; ws["A2"].font = F(italic=True, color="666666")
def hdr(ws, r, vals):
    for i, v in enumerate(vals):
        c = ws.cell(row=r, column=i + 1, value=v); c.font = H; c.fill = GREY
        c.alignment = Alignment(wrap_text=True, vertical="top")
def inp(ws, r, c, v):
    x = ws.cell(row=r, column=c, value=v); x.font = INP; x.fill = YEL; return x
def lab(ws, r, c, v, bold=False):
    x = ws.cell(row=r, column=c, value=v); x.font = F(bold=bold); return x
def fml(ws, r, c, f, fmt=None):
    x = ws.cell(row=r, column=c, value=f); x.font = F(bold=True)
    if fmt: x.number_format = fmt
    return x
def widths(ws, w):
    for i, x in enumerate(w): ws.column_dimensions[get_column_letter(i + 1)].width = x

# ---- Unit prices -------------------------------------------------------------
ws = wb.active; ws.title = "Unit prices"
title(ws, "The Tradeoff Method for ML Systems: cost-and-capacity toolkit",
      "Unit prices and per-accelerator capacities. Blue cells on yellow are inputs; edit them when prices move (checked September 2026). Every other tab reads from here. The book prints quantities and ratios only; this tab is where the money is.")
hdr(ws, 4, ["Name", "Item", "Value", "Unit", "Source / note"])
rows = [
 ("CORE_H", "Core (vCPU)", P["compute"]["core_usd_per_hour"], "$ per hour", "public list, Sept 2026"),
 ("ACC_RENT_H", "Accelerator, rented", P["compute"]["accelerator_rented_usd_per_hour"], "$ per hour", "mid-range data-centre accelerator, on demand"),
 ("ACC_OWN_H", "Accelerator, owned, amortised at full use", P["compute"]["accelerator_owned_usd_per_hour_at_full_use"], "$ per hour", "purchase + hosting over 3 years"),
 ("ACC_RATIO", "Accelerator-hour to core-hour", P["compute"]["accelerator_to_core_ratio"], "ratio", "Appendix A.9: 20–60×"),
 ("ACC_MEM", "Accelerator memory", P["accelerator"]["memory_gb"], "GB", "Appendix A.3: 40–200"),
 ("ACC_BW", "Accelerator memory bandwidth", P["accelerator"]["bandwidth_tb_per_s"], "TB/s", "Appendix A.3: 1.5–8"),
 ("PREFILL", "Prefill throughput, mid-sized model", P["accelerator"]["prefill_tokens_per_s_mid_model"], "tokens/s per accelerator", "Appendix A.3: 20k–100k"),
 ("DECODE", "Decode throughput, mid-sized model, batched", P["accelerator"]["decode_tokens_per_s_mid_model_batched"], "tokens/s per accelerator", "Appendix A.3: 2k–10k"),
 ("KV", "Key–value state per token, mid-sized model", P["accelerator"]["kv_bytes_per_token_mid_model"], "bytes", "Appendix A.3: 50–200 KB"),
 ("RANKER_CORE", "Small ranker, one item, on a core", P["per_item_cost"]["small_ranker_core_us"], "µs", "Appendix A.2: 20–50"),
 ("RANKER_ACC", "Small ranker, one item, batched on an accelerator", P["per_item_cost"]["small_ranker_accel_batched_us"], "µs", "Appendix A.2: 1–2"),
 ("EMB_ACC", "Text embedding, one passage, batched", P["per_item_cost"]["text_embedding_accel_batched_ms"], "ms", "Appendix A.2: 0.3–1"),
 ("IMG_ACC", "Image classifier, one image, batched", P["per_item_cost"]["image_classifier_accel_batched_ms"], "ms", "Appendix A.2: 1–3"),
 ("ANN", "Nearest-neighbour query over 100M vectors", P["per_item_cost"]["ann_query_100m_ms"], "ms", "Appendix A.1: 1–3"),
 ("TOK_IN", "Input tokens, rented", P["tokens"]["input_usd_per_million"], "$ per million", "hosted mid-sized model"),
 ("TOK_IN_C", "Input tokens, cached prefix", P["tokens"]["input_cached_usd_per_million"], "$ per million", ""),
 ("TOK_OUT", "Output tokens, rented", P["tokens"]["output_usd_per_million"], "$ per million", ""),
 ("OBJ", "Object storage", P["storage"]["object_usd_per_tb_month"], "$ per TB-month", ""),
 ("MEM", "Memory in a serving node", P["storage"]["memory_usd_per_gb_month"], "$ per GB-month", ""),
 ("GRAPH_X", "Graph index memory, multiple of raw vectors", P["vectors"]["graph_index_multiple"], "ratio", "Appendix A.5: 2–3×"),
 ("CLUSTER_X", "Cluster index memory, multiple of raw vectors", P["vectors"]["cluster_index_multiple"], "ratio", "Appendix A.5: ~1.1×"),
 ("JUDGE", "One human judgement", P["labels"]["human_judgement_usd_each"], "$", "Chapter 5: cents to dollars"),
 ("REV_DAY", "Reviewer decisions per day", P["labels"]["reviewer_decisions_per_day"], "per person", "Chapter 12"),
 ("REV_H", "Reviewer, loaded", P["labels"]["reviewer_usd_per_hour"], "$ per hour", "Chapter 13, Exercise 2"),
 ("DAY_S", "Seconds in the book's day", 100000, "s", "10^5"),
]
names = {}
for i, (n, item, v, unit, note) in enumerate(rows):
    r = 5 + i
    lab(ws, r, 1, n); lab(ws, r, 2, item); inp(ws, r, 3, v); lab(ws, r, 4, unit); lab(ws, r, 5, note)
    names[n] = f"'Unit prices'!$C${r}"
widths(ws, [14, 46, 12, 24, 40])

# ---- Fleet -------------------------------------------------------------------
ws = wb.create_sheet("Fleet from a rate")
title(ws, "Fleet from a rate (Chapters 4, 7)", "Items scored per second × cost per item = core- or accelerator-seconds per second; divide by the target utilisation. Example: the recommender's ranker, 1,000 req/s × 2,000 candidates.")
hdr(ws, 4, ["Input", "Value", "", "Result", "Value"])
lab(ws, 5, 1, "Requests per second at peak"); inp(ws, 5, 2, 1000)
lab(ws, 6, 1, "Items scored per request"); inp(ws, 6, 2, 2000)
lab(ws, 7, 1, "Cost per item on a core (µs)"); inp(ws, 7, 2, 30)
lab(ws, 8, 1, "Cost per item batched on an accelerator (µs)"); inp(ws, 8, 2, 2)
lab(ws, 9, 1, "Target utilisation"); inp(ws, 9, 2, 0.7)
lab(ws, 5, 4, "Items per second"); fml(ws, 5, 5, "=B5*B6", "#,##0")
lab(ws, 6, 4, "Core-seconds per second"); fml(ws, 6, 5, "=E5*B7/1000000", "0.0")
lab(ws, 7, 4, "Cores at utilisation"); fml(ws, 7, 5, "=ROUNDUP(E6/B9,0)", "#,##0")
lab(ws, 8, 4, "Accelerator-seconds per second"); fml(ws, 8, 5, "=E5*B8/1000000", "0.0")
lab(ws, 9, 4, "Accelerators at utilisation"); fml(ws, 9, 5, "=ROUNDUP(E8/B9,0)", "#,##0")
lab(ws, 10, 4, "Core-seconds per thousand requests"); fml(ws, 10, 5, "=B6*B7/1000000*1000", "0.0")
lab(ws, 11, 4, "Cores, $ per hour at utilisation"); fml(ws, 11, 5, f"=E6/B9*{names['CORE_H']}", "$#,##0.00")
lab(ws, 12, 4, "Accelerators, $ per hour rented"); fml(ws, 12, 5, f"=E9*{names['ACC_RENT_H']}", "$#,##0.00")
widths(ws, [44, 12, 3, 40, 14])

# ---- Generative serving --------------------------------------------------------
ws = wb.create_sheet("Generative serving")
title(ws, "Generative serving (Chapters 8, 13)", "Prefill is compute-bound; decode is memory-bound and the in-flight state caps the batch. Size the two pools separately and take the larger of decode-by-throughput and decode-by-memory. Example: the serving platform, 2,000 req/s.")
hdr(ws, 4, ["Input", "Value", "", "Result", "Value"])
ins = [("Requests per second at peak", 2000), ("Input tokens per request", 4000), ("Cached shared-prefix tokens", 3000),
       ("Output tokens per request", 300), ("Seconds a request stays in flight", 6), ("Model parameters (billions)", 10),
       ("Bytes per parameter (2 half, 1 8-bit)", 2)]
for i, (l, v) in enumerate(ins): lab(ws, 5 + i, 1, l); inp(ws, 5 + i, 2, v)
res = [("Prefill tokens per second, no cache", "=B5*B6", "#,##0"),
       ("Prefill accelerators, no cache", f"=ROUNDUP(E5/{names['PREFILL']},0)", "#,##0"),
       ("Prefill tokens per second, cached", "=B5*(B6-B7)", "#,##0"),
       ("Prefill accelerators, cached", f"=ROUNDUP(E7/{names['PREFILL']},0)", "#,##0"),
       ("Decode tokens per second", "=B5*B8", "#,##0"),
       ("Decode accelerators by throughput", f"=ROUNDUP(E9/{names['DECODE']},0)", "#,##0"),
       ("Model bytes (GB)", "=B10*B11", "0.0"),
       ("Free memory per accelerator (GB)", f"={names['ACC_MEM']}-E11", "0.0"),
       ("KV state per request (GB)", f"=(B6+B8)*{names['KV']}/1000000000", "0.00"),
       ("Sequences in flight per accelerator", "=INT(E12/E13)", "#,##0"),
       ("Sequences in flight across the pool", "=B5*B9", "#,##0"),
       ("Decode accelerators by memory", "=ROUNDUP(E15/E14,0)", "#,##0"),
       ("Decode pool (the larger)", "=MAX(E10,E16)", "#,##0"),
       ("Total at peak, cached prefix", "=E8+E17", "#,##0"),
       ("Single-sequence decode rate (tokens/s)", f"={names['ACC_BW']}*1000000000000/(E11*1000000000)", "#,##0"),
       ("Accelerator-seconds per request", f"=(B6-B7)/{names['PREFILL']}+B8/{names['DECODE']}", "0.000"),
       ("Rented $ per thousand requests", f"=((B6-B7)*{names['TOK_IN']}+B7*{names['TOK_IN_C']}+B8*{names['TOK_OUT']})/1000000*1000", "$#,##0.00"),
       ("Output tokens' share of the rented bill", f"=(B8*{names['TOK_OUT']})/((B6-B7)*{names['TOK_IN']}+B7*{names['TOK_IN_C']}+B8*{names['TOK_OUT']})", "0%")]
for i, (l, f, fmt) in enumerate(res): lab(ws, 5 + i, 4, l); fml(ws, 5 + i, 5, f, fmt)
widths(ws, [40, 12, 3, 44, 16])

# ---- Own against rent ----------------------------------------------------------
ws = wb.create_sheet("Own vs rent")
title(ws, "Own against rent (Chapter 8)", "Owning is priced for peak and paid whether used or not; renting is paid for what is used. The break-even is a utilisation. Examples: the support assistant (2 accelerators at 7%) and the platform (160 at 40–70%).")
hdr(ws, 4, ["Input", "Value", "", "Result", "Value"])
lab(ws, 5, 1, "Accelerators at peak"); inp(ws, 5, 2, 160)
lab(ws, 6, 1, "Average utilisation over the day"); inp(ws, 6, 2, 0.4)
lab(ws, 5, 4, "Own, $ per year"); fml(ws, 5, 5, f"=B5*8760*{names['ACC_OWN_H']}", "$#,##0")
lab(ws, 6, 4, "Rent, $ per year"); fml(ws, 6, 5, f"=B5*B6*8760*{names['ACC_RENT_H']}", "$#,##0")
lab(ws, 7, 4, "Break-even utilisation"); fml(ws, 7, 5, f"={names['ACC_OWN_H']}/{names['ACC_RENT_H']}", "0%")
lab(ws, 8, 4, "Decision"); fml(ws, 8, 5, '=IF(E5<E6,"own","rent")')
widths(ws, [40, 12, 3, 30, 16])

# ---- Vectors and indexes -------------------------------------------------------
ws = wb.create_sheet("Index and false matches")
title(ws, "Vectors, indexes and the false-match budget (Chapters 7, 12)", "Memory for N vectors at d dims and bytes per dim, under a graph or cluster index; and the per-pair false-match probability a threshold must deliver for a ceiling of wrongful matches a day, tested over K candidates. Example: near-duplicate search, 2B vectors.")
hdr(ws, 4, ["Input", "Value", "", "Result", "Value"])
ins = [("Vectors", 2000000000), ("Dimensions", 128), ("Bytes per dimension (4 fp32, 1 int8, 0.25 for 32-byte codes at 128 dims)", 0.25),
       ("Uploads (queries) per day", 5000000), ("Candidates the threshold is applied to", 20), ("Ceiling on wrongful matches per day", 100)]
for i, (l, v) in enumerate(ins): lab(ws, 5 + i, 1, l); inp(ws, 5 + i, 2, v)
res = [("Raw vectors (GB)", "=B5*B6*B7/1000000000", "#,##0"),
       ("Graph index (GB)", f"=E5*{names['GRAPH_X']}", "#,##0"),
       ("Cluster index (GB)", f"=E5*{names['CLUSTER_X']}", "#,##0"),
       ("Cluster index, $ per month in memory", f"=E7*{names['MEM']}", "$#,##0"),
       ("Per-pair false-match probability required", "=B10/(B8*B9)", "0.0E+00"),
       ("Same, if applied over every vector", "=B10/(B8*B5)", "0.0E+00")]
for i, (l, f, fmt) in enumerate(res): lab(ws, 5 + i, 4, l); fml(ws, 5 + i, 5, f, fmt)
widths(ws, [64, 14, 3, 40, 16])

# ---- Labels, logs and cadence --------------------------------------------------
ws = wb.create_sheet("Labels, logs, cadence")
title(ws, "Labels, logs and cadence (Chapters 5, 6, 9, 12)", "The served-feature log sized from the label delay; a retrain cadence priced; a judged set sized from the difference to detect; a prevalence sample sized from the precision wanted.")
hdr(ws, 4, ["Input", "Value", "", "Result", "Value"])
ins = [("Items served per second", 10000), ("Bytes per served-feature row", 800), ("Label delay (days)", 90), ("Training window (days)", 90),
       ("Retrain: hours per run", 300), ("Retrain: runs per day", 1), ("Retrain on accelerators? (1 yes, 0 no)", 0),
       ("Judged set: rate being measured p", 0.9), ("Judged set: difference to detect", 0.02),
       ("Prevalence being measured", 0.0005), ("Relative half-width wanted", 0.2)]
for i, (l, v) in enumerate(ins): lab(ws, 5 + i, 1, l); inp(ws, 5 + i, 2, v)
res = [("Log retention (days)", "=B7+B8", "#,##0"),
       ("Log, GB per day", f"=B5*{names['DAY_S']}*B6/1000000000", "#,##0"),
       ("Log, TB total", "=E6*E5/1000", "#,##0"),
       ("Log, $ per month in object storage", f"=E7*{names['OBJ']}", "$#,##0"),
       ("Retrain, $ per day", f"=B9*B10*IF(B11=1,{names['ACC_RENT_H']},{names['CORE_H']})", "$#,##0"),
       ("Retrain, $ per year", "=E9*365", "$#,##0"),
       ("Judged set size n", "=ROUNDUP(4*B12*(1-B12)/B13^2,0)", "#,##0"),
       ("Judged set, $ at the human rate", f"=E11*{names['JUDGE']}", "$#,##0"),
       ("Prevalence sample, views per week", "=ROUNDUP(1.96^2*(1-B14)/(B15^2*B14),0)", "#,##0"),
       ("Reviewers, full time", f"=E13/7/{names['REV_DAY']}", "0.0")]
for i, (l, f, fmt) in enumerate(res): lab(ws, 5 + i, 4, l); fml(ws, 5 + i, 5, f, fmt)
widths(ws, [44, 14, 3, 40, 16])

wb.save("ml_capacity_toolkit.xlsx")
print("wrote ml_capacity_toolkit.xlsx")
